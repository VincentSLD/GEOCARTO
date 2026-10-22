# -*- coding: utf-8 -*-
"""Génère le dossier de commercialisation SaaS de GEOCARTO au format .docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ORANGE = RGBColor(0xE6, 0x51, 0x00)
GREY_FILL = "F3F3F3"
CODE_FILL = "1E1E1E"

doc = Document()

# ---- Styles de base ----
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
for name, size in [('Title', 26), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)]:
    st = doc.styles[name]
    st.font.color.rgb = ORANGE
    st.font.size = Pt(size)
    st.font.bold = True

def set_cell_bg(cell, fill):
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def add_code(text, dark=True):
    """Bloc de code monospace dans un tableau 1x1 ombré."""
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, CODE_FILL if dark else GREY_FILL)
    cell.paragraphs[0].text = ''
    first = True
    for line in text.rstrip('\n').split('\n'):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line if line != '' else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xE8, 0xE8, 0xE8) if dark else RGBColor(0x22, 0x22, 0x22)
    doc.add_paragraph()

def add_note(text, label="À RETENIR"):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, "FFF3E9")
    p = cell.paragraphs[0]
    r = p.add_run(f"⚠ {label} — ")
    r.bold = True; r.font.color.rgb = ORANGE; r.font.size = Pt(10.5)
    r2 = p.add_run(text); r2.font.size = Pt(10.5)
    doc.add_paragraph()

def bullets(items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        if isinstance(it, tuple):
            r = p.add_run(it[0] + " : "); r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)

def numbered(items):
    for it in items:
        p = doc.add_paragraph(style='List Number')
        if isinstance(it, tuple):
            r = p.add_run(it[0] + " — "); r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].paragraphs[0].text = ''
        run = hdr[i].paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].text = ''
            run = cells[i].paragraphs[0].add_run(str(val)); run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()

def add_toc():
    p = doc.add_paragraph()
    run = p.add_run()
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "Ctrl+clic ici puis « Mettre à jour les champs » (F9) pour générer le sommaire."
    fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
    for el in (fld1, instr, fld2, t, fld3):
        run._r.append(el)

# ================= PAGE DE GARDE =================
for _ in range(3):
    doc.add_paragraph()
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run("GEOCARTO"); r.bold = True; r.font.size = Pt(40); r.font.color.rgb = ORANGE
sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sp.add_run("Dossier de commercialisation en mode SaaS"); r.font.size = Pt(18); r.bold = True
sp2 = doc.add_paragraph(); sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sp2.add_run("Architecture, jetons d'utilisation, paiement en ligne,\nespace client et feuille de route"); r.font.size = Pt(12); r.italic = True
for _ in range(6):
    doc.add_paragraph()
meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Version 1.0 — 16 septembre 2026"); r.font.size = Pt(11); r.bold = True
meta2 = doc.add_paragraph(); meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta2.add_run("Document de travail interne — confidentiel"); r.font.size = Pt(10); r.italic = True; r.font.color.rgb = RGBColor(0x88,0x88,0x88)
doc.add_page_break()

# ================= SOMMAIRE =================
doc.add_heading("Sommaire", level=1)
add_toc()
doc.add_page_break()

# ================= 1. RESUME EXECUTIF =================
doc.add_heading("1. Résumé exécutif", level=1)
doc.add_paragraph(
    "GEOCARTO est aujourd'hui un outil cartographique de recherche de prestataires (pelles mécaniques) "
    "permettant de localiser un point d'intervention, de trouver les prestataires couvrant la zone, de "
    "calculer les temps de trajet routiers réels et d'afficher les itinéraires. L'objectif de ce dossier "
    "est de définir comment transformer cet outil en produit commercialisable, facturé à l'usage via un "
    "système de jetons, avec espace client, paiement en ligne, suivi de crédits et sauvegarde des dossiers "
    "d'investigation."
)
doc.add_paragraph(
    "La bonne nouvelle : la stack actuelle (front statique sur Vercel + Supabase) constitue déjà une "
    "excellente fondation. Il n'est pas nécessaire de tout réécrire. Le chantier principal consiste à "
    "déplacer la logique facturée côté serveur, à ajouter une isolation multi-locataires (multi-tenant), "
    "à brancher un système de paiement (Stripe) et à professionnaliser le service de routage."
)
add_note(
    "Le point le plus important de tout le projet : aucune requête facturée ne doit être décomptée ou "
    "exécutée uniquement côté navigateur. Le décompte de jetons et l'appel aux services payants doivent "
    "impérativement passer par un serveur que vous seul contrôlez.",
    label="POINT CRITIQUE")

# ================= 2. CONTEXTE & EXISTANT =================
doc.add_heading("2. Contexte & analyse de l'existant", level=1)
doc.add_paragraph("La solution actuelle repose sur :")
bullets([
    ("Front-end", "pages HTML/JS statiques (ex. pelles.html) hébergées sur Vercel, cartographie Leaflet."),
    ("Base de données", "Supabase (PostgreSQL) — tables des prestataires, matériels, tarifs par zone."),
    ("Routage", "serveur de démonstration public OSRM (router.project-osrm.org)."),
    ("Données externes", "API Recherche d'entreprises et API Adresse (data.gouv.fr) pour le SIRET et le géocodage."),
])
doc.add_paragraph("Limites actuelles vis-à-vis d'une commercialisation :")
bullets([
    "Toute la logique s'exécute dans le navigateur : impossible de facturer de façon fiable.",
    "La clé Supabase est publique (anon) : convient à un usage interne, pas à un multi-client cloisonné.",
    "Le serveur de routage est une démo publique : ni SLA, ni droit d'usage commercial, ni garantie de disponibilité.",
    "Aucune notion de compte client, de crédit, de facturation ou de sauvegarde de dossier.",
])

# ================= 3. PRINCIPE BACKEND-FIRST =================
doc.add_heading("3. Principe fondamental : passer en « backend-first »", level=1)
doc.add_paragraph(
    "Pour un produit payant à la requête, chaque action facturée doit transiter par une couche serveur "
    "maîtrisée. Le navigateur ne doit plus jamais appeler directement les services payants (routage, "
    "géocodage) ni gérer lui-même le solde de crédits."
)
doc.add_paragraph("Le déroulé sécurisé d'une requête facturée :")
numbered([
    "Le front envoie la demande (ex. « recherche autour de ce point ») à une fonction serveur, avec le jeton d'authentification de l'utilisateur.",
    "Le serveur vérifie l'identité et l'organisation de l'utilisateur.",
    "Le serveur débite les crédits de façon atomique (refus si solde insuffisant → invitation à recharger).",
    "Le serveur appelle le service de routage / géocodage.",
    "Le serveur renvoie le résultat ET le nouveau solde au front.",
])
add_note(
    "Si le décompte se fait côté client, n'importe quel utilisateur peut l'annuler depuis la console du "
    "navigateur et consommer le service gratuitement. C'est la raison n°1 de la bascule backend.")

# ================= 4. ARCHITECTURE CIBLE =================
doc.add_heading("4. Architecture cible", level=1)
doc.add_paragraph("Schéma logique des flux :")
add_code(
"""Navigateur (front Vercel)
        │  1. requête + JWT
        ▼
Fonction serveur (Supabase Edge Function / Vercel Serverless)
        │  2. vérifie l'utilisateur (Supabase Auth)
        │  3. consume_credits()  ── atomique ──►  PostgreSQL (Supabase)
        │        └── si solde insuffisant : HTTP 402 + lien d'achat
        │  4. appel routage/géocodage
        ▼
Service de routage (OSRM auto-hébergé ou fournisseur commercial)
        │
        ▼
        5. résultat + nouveau solde ──► Navigateur

Paiement :  Stripe Checkout ──► Webhook ──► crédite le ledger (+N jetons)
""", dark=True)

# ================= 5. STACK TECHNIQUE =================
doc.add_heading("5. Stack technique recommandée", level=1)
doc.add_paragraph(
    "La recommandation capitalise au maximum sur les briques déjà en place afin de limiter le coût "
    "d'apprentissage et de migration.")
table(
    ["Besoin", "Solution recommandée", "Pourquoi"],
    [
        ["Hébergement front", "Vercel (déjà en place)", "Aucun changement nécessaire"],
        ["Auth + base + sécurité", "Supabase (déjà en place)", "Auth, PostgreSQL, Row-Level Security, Edge Functions"],
        ["Logique serveur / proxy", "Supabase Edge Functions ou Vercel Serverless", "Emplacement du décompte de jetons"],
        ["Paiement", "Stripe", "Standard, gère la TVA et les factures UE"],
        ["Routage (production)", "OSRM/Valhalla auto-hébergé ou Geoapify/Mapbox/ORS", "Le serveur démo public est interdit en prod"],
        ["Emails transactionnels", "Resend ou Postmark", "Reçus, alertes de crédits bas, onboarding"],
        ["Supervision", "Sentry + logs Supabase/Vercel", "Détection des erreurs et abus"],
    ],
    widths=[1.6, 2.4, 2.6])

# ================= 6. ESPACE CLIENT / MULTI-TENANT =================
doc.add_heading("6. Espace client & architecture multi-locataires", level=1)
doc.add_paragraph(
    "Un SaaS professionnel doit isoler strictement les données de chaque client. On raisonne en "
    "« organisations » plutôt qu'en simples utilisateurs, car un client est souvent une entreprise avec "
    "plusieurs collaborateurs partageant le même crédit.")
doc.add_heading("6.1 Authentification", level=2)
bullets([
    "Supabase Auth : email + mot de passe, lien magique, ou connexion Google.",
    "Chaque utilisateur appartient à une ou plusieurs organisations, avec un rôle (propriétaire, admin, membre).",
])
doc.add_heading("6.2 Isolation par Row-Level Security (RLS)", level=2)
doc.add_paragraph(
    "La RLS de PostgreSQL/Supabase garantit qu'un client ne peut lire ou modifier que les données de son "
    "organisation, directement au niveau de la base — c'est la pièce maîtresse de la sécurité multi-tenant.")
add_code(
"""-- Tables de base
organizations (id, nom, tva_intra, credits_balance, plan, created_at)
memberships   (org_id, user_id, role)   -- role: owner | admin | member

-- Exemple de policy RLS : un membre ne voit que SES organisations
create policy "membres voient leur org"
  on organizations for select
  using ( id in (select org_id from memberships where user_id = auth.uid()) );
""", dark=False)

# ================= 7. SYSTEME DE JETONS =================
doc.add_heading("7. Système de jetons / crédits", level=1)
doc.add_paragraph("C'est le cœur du modèle économique. Deux principes : traçabilité et atomicité.")
doc.add_heading("7.1 Modèle de données", level=2)
add_code(
"""credit_ledger (id, org_id, delta, reason, ref, created_at)
   -- +100 « achat pack », -1 « recherche », -1 « export »…
   -- Le solde = SOMME(delta). C'est la source de vérité auditable.

organizations.credits_balance  -- cache du solde pour lecture rapide
""", dark=False)
doc.add_heading("7.2 Décompte atomique (indispensable)", level=2)
doc.add_paragraph(
    "Deux requêtes simultanées ne doivent jamais pouvoir dépenser le même crédit. On utilise une fonction "
    "PostgreSQL qui décrémente et journalise en une seule opération, appelée par la fonction serveur :")
add_code(
"""create function consume_credits(p_org uuid, p_amount int, p_reason text)
returns int language plpgsql security definer as $$
declare new_balance int;
begin
  update organizations set credits_balance = credits_balance - p_amount
   where id = p_org and credits_balance >= p_amount
   returning credits_balance into new_balance;
  if not found then raise exception 'INSUFFICIENT_CREDITS'; end if;
  insert into credit_ledger(org_id, delta, reason)
       values (p_org, -p_amount, p_reason);
  return new_balance;
end $$;
""", dark=True)
doc.add_heading("7.3 Que vaut « une requête » ?", level=2)
doc.add_paragraph(
    "Une recherche « point d'intervention » déclenche aujourd'hui jusqu'à 20 appels de routage plus du "
    "géocodage. Deux façons de tarifer :")
bullets([
    ("Forfait par action (recommandé)", "1 crédit = 1 recherche, quel que soit le nombre d'itinéraires calculés en coulisse. Simple à comprendre pour le client ; vous absorbez la complexité technique dans votre marge."),
    ("Pondéré", "recherche simple = 1, recherche avec itinéraires routiers = 3, export = 1, etc. Plus juste mais moins lisible."),
])

# ================= 8. PAIEMENT EN LIGNE =================
doc.add_heading("8. Paiement en ligne (Stripe)", level=1)
bullets([
    ("Packs de jetons (paiement unique)", "via Stripe Checkout — ex. 100 crédits = 29 €, 500 = 99 €. Idéal pour démarrer."),
    ("Abonnement", "quota mensuel inclus + dépassement facturé, via Stripe Billing. À introduire plus tard."),
])
add_note(
    "Le crédit ne doit être ajouté QUE par le webhook Stripe (événement « checkout.session.completed »), "
    "jamais par le front. Le webhook insère « +N » dans le ledger. Cela évite toute fraude au rechargement.")
doc.add_paragraph("Stripe prend en charge nativement :")
bullets([
    "la TVA intracommunautaire et les taux par pays ;",
    "la génération des factures et reçus PDF ;",
    "les moyens de paiement (CB, virement SEPA, etc.).",
])
doc.add_paragraph("Points juridiques à cadrer : CGV/CGU, politique des crédits (remboursables ? expirent-ils ?), mentions légales.")

# ================= 9. ROUTAGE & DONNEES =================
doc.add_heading("9. Routage & données : sortir des services « démo »", level=1)
doc.add_paragraph(
    "Le serveur router.project-osrm.org est une démonstration : son usage commercial est interdit et il "
    "n'offre aucune garantie. Deux stratégies pour la production :")
table(
    ["Option", "Avantages", "Inconvénients"],
    [
        ["OSRM / Valhalla auto-hébergé (VPS)", "Coût fixe, volume illimité, données maîtrisées", "Maintenance serveur, mise à jour des cartes"],
        ["Fournisseur commercial (Geoapify, Mapbox, ORS…)", "Zéro maintenance, mise en route rapide", "Coût variable à répercuter dans le prix du jeton"],
    ],
    widths=[2.3, 2.3, 2.3])
add_note(
    "Votre véritable actif, c'est la base de prestataires (contacts, tarifs, rayons d'action). Ne jamais "
    "l'exposer en bloc au client : ne servir que les résultats d'une recherche, limiter le débit "
    "(rate-limiting) et surveiller les tentatives d'extraction massive (scraping).", label="ACTIF STRATÉGIQUE")

# ================= 10. INVESTIGATIONS =================
doc.add_heading("10. Sauvegarde des points d'investigation", level=1)
doc.add_paragraph(
    "Permettre aux clients d'enregistrer leurs points d'investigation avec références transforme l'outil "
    "en mini-CRM — un argument de vente fort.")
add_code(
"""investigations (id, org_id, created_by, titre, reference_chantier,
                client_final, adresse, lat, lng, notes, statut, created_at)
   -- statut : à chiffrer | envoyé | gagné | perdu

investigation_results (investigation_id, prestataire_id,
                       distance_m, duree_s, tarif, snapshot jsonb)
   -- snapshot : fige le résultat (tarif/temps) au moment de la recherche
""", dark=False)
bullets([
    ("Snapshot des résultats", "les tarifs prestataires évolueront ; le devis envoyé au client doit rester traçable dans le temps."),
    ("Champs métier", "référence chantier, client final, statut d'avancement → suivi commercial intégré."),
    ("Export", "génération PDF/CSV d'un dossier complet pour transmission ou archivage."),
])

# ================= 11. SECURITE & RGPD =================
doc.add_heading("11. Sécurité & conformité RGPD", level=1)
add_note(
    "Les contacts prestataires sont des données personnelles professionnelles. Vendre l'accès à ces "
    "données impose un cadre RGPD rigoureux — incontournable pour commercialiser en France et dans l'UE.",
    label="OBLIGATION LÉGALE")
bullets([
    "Base légale du traitement et information des personnes concernées.",
    "CGU / CGV, mentions légales, politique de confidentialité.",
    "Accords de sous-traitance (DPA) avec Supabase, Stripe, le fournisseur de routage.",
    "Droit d'accès, de rectification et de suppression des données.",
    "Sécurité : chiffrement, rotation des clés, journaux d'audit, limitation de débit (rate-limiting), détection d'usage anormal.",
])

# ================= 12. BACK-OFFICE =================
doc.add_heading("12. Back-office administrateur (côté éditeur)", level=1)
doc.add_paragraph("Interface d'administration réservée à vous, l'éditeur, pour piloter le service :")
bullets([
    "Vue des clients, de leur usage et de leur solde de crédits.",
    "Créditer manuellement une organisation (geste commercial, SAV).",
    "Suivi du chiffre d'affaires et des consommations.",
    "Gestion et mise à jour de la base de prestataires.",
    "Consultation des journaux d'audit et alertes de sécurité.",
])

# ================= 13. FONCTIONS COMPLEMENTAIRES =================
doc.add_heading("13. Fonctionnalités complémentaires à prévoir", level=1)
bullets([
    ("Rôles & équipes", "plusieurs utilisateurs par client, crédit partagé, gestion des droits."),
    ("Emails transactionnels", "reçu d'achat, alerte « crédits bas », message de bienvenue / onboarding."),
    ("Tableau de bord client", "solde, historique de consommation, factures téléchargeables."),
    ("Supervision", "logs, alertes (Sentry), monitoring de la disponibilité du routage."),
    ("Essai gratuit", "quelques crédits offerts à l'inscription pour tester sans friction."),
    ("Sauvegardes & résilience", "de la base de données, votre actif principal."),
    ("Support & documentation", "aide en ligne, contact, FAQ."),
    ("API programmatique (plus tard)", "accès versionné pour intégrations tierces, source de revenus additionnelle."),
])

# ================= 14. MODELE ECONOMIQUE =================
doc.add_heading("14. Modèle économique & tarification des jetons", level=1)
doc.add_paragraph(
    "Principe : calculer le coût réel d'une recherche (routage + géocodage + hébergement), puis appliquer "
    "une marge. Un ordre de grandeur indicatif :")
table(
    ["Élément", "Ordre de grandeur", "Commentaire"],
    [
        ["Coût infra par recherche", "≈ 0,01 – 0,03 €", "Routage + géocodage + hébergement"],
        ["Prix de vente du crédit", "0,20 – 0,50 €", "Marge confortable, indolore pour un pro"],
        ["Palier d'essai gratuit", "5 – 10 crédits", "Réduit la friction à l'inscription"],
        ["Pack de démarrage", "≈ 100 crédits / 29 €", "Point d'entrée typique"],
    ],
    widths=[2.2, 1.8, 2.5])
doc.add_paragraph(
    "Justification : un professionnel qui décroche un chantier de plusieurs milliers d'euros grâce à une "
    "recherche ne sera pas sensible à un coût de quelques dizaines de centimes par requête.")

# ================= 15. FEUILLE DE ROUTE =================
doc.add_heading("15. Feuille de route (MVP par phases)", level=1)
doc.add_paragraph("Ne pas tout construire d'un coup. Ordre de priorité recommandé :")
table(
    ["Phase", "Contenu", "Objectif"],
    [
        ["1", "Auth + organisations + RLS", "Isoler les données par client"],
        ["2", "Fonction serveur proxy + décompte de crédits", "Facturation fiable à l'usage"],
        ["3", "Stripe Checkout (packs) + webhook de crédit", "Encaisser en ligne"],
        ["4", "Sauvegarde des investigations + historique", "Fidélisation, valeur métier"],
        ["5", "Back-office administrateur minimal", "Piloter le service"],
        ["6", "Abonnements, équipes, PDF, alertes, RGPD affiné", "Industrialisation"],
    ],
    widths=[0.7, 3.4, 2.4])

# ================= 16. COUTS INFRA =================
doc.add_heading("16. Estimation des coûts d'infrastructure", level=1)
doc.add_paragraph("Ordres de grandeur mensuels pour un démarrage (à ajuster selon le volume) :")
table(
    ["Poste", "Coût mensuel indicatif", "Note"],
    [
        ["Vercel", "0 – 20 €", "Gratuit possible au début"],
        ["Supabase", "0 – 25 €", "Palier gratuit puis Pro"],
        ["Stripe", "Commission ~1,5 % + 0,25 €/transaction", "Pas d'abonnement fixe"],
        ["Routage auto-hébergé (VPS)", "5 – 20 €", "Un petit serveur suffit au début"],
        ["Emails (Resend/Postmark)", "0 – 15 €", "Selon volume"],
        ["Supervision (Sentry)", "0 – 26 €", "Palier gratuit au début"],
    ],
    widths=[2.2, 2.3, 2.0])
doc.add_paragraph("Total de départ réaliste : quelques dizaines d'euros par mois, qui croissent avec l'usage — modèle sain pour un SaaS.")

# ================= 17. RISQUES =================
doc.add_heading("17. Risques & points de vigilance", level=1)
bullets([
    ("Conformité RGPD", "principal risque juridique — à traiter dès le départ."),
    ("Dépendance au routage", "prévoir un plan B (bascule de fournisseur) ; la logique est isolée dans une seule fonction pour faciliter le changement."),
    ("Protection de la base", "éviter le scraping et la revente non autorisée de vos données."),
    ("Fraude au paiement", "s'appuyer strictement sur les webhooks Stripe, jamais sur le front."),
    ("Sur-ingénierie", "livrer le MVP par phases plutôt que viser la solution parfaite d'emblée."),
])

# ================= 18. PROCHAINES ETAPES =================
doc.add_heading("18. Prochaines étapes concrètes", level=1)
numbered([
    "Valider le périmètre du MVP (phases 1 à 3).",
    "Rédiger le schéma SQL complet (organisations, membres, ledger, investigations) avec les policies RLS.",
    "Développer la première fonction serveur « recherche » qui débite un crédit et proxifie le routage.",
    "Mettre en place Stripe Checkout et le webhook de crédit.",
    "Choisir la solution de routage de production (auto-hébergé vs commercial) et migrer.",
    "Cadrer le volet juridique (RGPD, CGV/CGU) avant l'ouverture commerciale.",
])
doc.add_paragraph()
closing = doc.add_paragraph()
r = closing.add_run(
    "En synthèse : pas de changement de technologie majeur. L'effort porte sur le déplacement de la "
    "logique facturée côté serveur, l'isolation multi-locataires (RLS), le branchement de Stripe et la "
    "professionnalisation du routage. Le reste (investigations, back-office, conformité) se construit "
    "progressivement par-dessus cette fondation.")
r.italic = True

doc.save(r"c:\PROJETS\01_GEOCARTO\GEOCARTO_Dossier_Commercialisation_SaaS.docx")
print("OK - document généré")
